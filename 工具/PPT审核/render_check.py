# -*- coding: utf-8 -*-
"""渲染级版面校验：把 PPTX 转 PDF，用 span 级实测每个文本片段的真实位置，
检查其是否越出所属卡片（取包含该片段中心的最小形状）或越出页面。
用法： python render_check.py <pptx路径>
依赖： LibreOffice(soffice)、pymupdf
"""
import os
import shutil
import subprocess
import sys
import tempfile

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def find_soffice():
    """定位 soffice：环境变量 → PATH → 常见安装位置。"""
    env = os.environ.get("SOFFICE")
    if env and os.path.exists(env):
        return env
    hit = shutil.which("soffice") or shutil.which("libreoffice")
    if hit:
        return hit
    for p in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
    ):
        if os.path.exists(p):
            return p
    return None


SOFFICE = find_soffice()
TOL = 1.5  # pt 容差


def containers(prs):
    out = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                continue
            if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
                continue
            L = sh.left.inches * 72
            T = sh.top.inches * 72
            R = L + sh.width.inches * 72
            B = T + sh.height.inches * 72
            out.append((si, L, T, R, B))
    return out


def spans(page):
    for blk in page.get_text('dict')['blocks']:
        for line in blk.get('lines', []):
            for sp in line.get('spans', []):
                t = (sp.get('text') or '').strip()
                if t:
                    yield sp['bbox'], t


def check(pptx):
    if not SOFFICE:
        return [('RENDER', '-', '未找到 LibreOffice(soffice)，请设环境变量 SOFFICE 指向可执行文件', '-')]
    tmp = tempfile.mkdtemp(prefix='render_check_')
    subprocess.run([SOFFICE, '--headless', '--convert-to', 'pdf', '--outdir', tmp, pptx],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.join(tmp, os.path.splitext(os.path.basename(pptx))[0] + '.pdf')
    if not os.path.exists(pdf):
        return [('RENDER', '-', '无法转 PDF', pdf)]

    prs = Presentation(pptx)
    conts = containers(prs)
    doc = pymupdf.open(pdf)
    issues = []
    for pi, page in enumerate(doc, 1):
        PW, PH = page.rect.width, page.rect.height
        for (x0, y0, x1, y1), txt in spans(page):
            if x0 < 1.5 or y0 < 1.5 or x1 > PW - 1.5 or y1 > PH - 1.5:
                issues.append(('PAGE', 'p%d' % pi, '文字越出页面',
                               'x1=%.0f y1=%.0f (页 %.0f×%.0f) | %s' % (x1, y1, PW, PH, txt[:24])))
            cx, cy = (x0 + x1) / 2.0, (y0 + y1) / 2.0
            best = None
            for (si, L, T, R, B) in conts:
                if si != pi:
                    continue
                if L <= cx <= R and T <= cy <= B:
                    area = (R - L) * (B - T)
                    if best is None or area < best[0]:
                        best = (area, L, T, R, B)
            if best is None:
                continue
            _, L, T, R, B = best
            if x1 > R - TOL or y1 > B - TOL or x0 < L + TOL or y0 < T + TOL:
                issues.append(('CARD', 'p%d' % pi,
                               '文字越出所属卡片 L=%.0f T=%.0f R=%.0f B=%.0f' % (L, T, R, B),
                               'x0=%.0f y0=%.0f x1=%.0f y1=%.0f | %s' % (x0, y0, x1, y1, txt[:26])))
    return issues


if __name__ == '__main__':
    items = check(sys.argv[1])
    if not items:
        print('PASS（渲染实测）：文字均在容器与页面内')
    else:
        print('FAIL：实测 %d 处越界' % len(items))
        for kind, pg, msg, extra in items:
            print('  [%s] %s %s | %s' % (kind, pg, msg, extra))
