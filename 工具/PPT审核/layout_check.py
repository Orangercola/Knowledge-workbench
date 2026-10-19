# -*- coding: utf-8 -*-
"""PPTX 版面溢出审核：按字符宽度估算文本排版高度，与文本框可用高度比对。
用法： python layout_check.py <pptx路径>
"""
import math
import sys
from pptx import Presentation


def is_fullwidth(ch):
    o = ord(ch)
    if o < 0x1100:
        return False
    return (
        (0x1100 <= o <= 0x115F) or
        (0x2E80 <= o <= 0xA4CF) or
        (0xAC00 <= o <= 0xD7A3) or
        (0xF900 <= o <= 0xFAFF) or
        (0xFE10 <= o <= 0xFE6F) or
        (0xFF00 <= o <= 0xFF60) or
        (0xFFE0 <= o <= 0xFFE6) or
        (0x1F300 <= o <= 0x1FAFF)
    )


def char_w(ch, size):
    if is_fullwidth(ch):
        return size
    if ch == ' ':
        return size * 0.28
    if ch in '·:：；，。、！？()（）「」《》“”‘’—–-+/%':
        return size * 0.55
    return size * 0.52


def line_spacing(p):
    ls = p.line_spacing
    if ls is None:
        return 1.2
    try:
        return float(ls)
    except (TypeError, ValueError):
        # Length 对象（绝对行距）
        try:
            return ls.pt
        except Exception:
            return 1.2


def para_height(p, width_pt, absolute_ls=False):
    runs = [(r.text or '', (r.font.size.pt if r.font.size else 18.0)) for r in p.runs]
    sb = p.space_before.pt if p.space_before is not None else 0.0
    sa = p.space_after.pt if p.space_after is not None else 0.0
    if not runs:
        return sb + sa
    size = max(s for _, s in runs)
    w = sum(char_w(c, s) for t, s in runs for c in t)
    lines = max(1, int(math.ceil(w / width_pt - 1e-6))) if width_pt > 0 else 1
    ls = line_spacing(p)
    lh = ls if absolute_ls else size * ls
    return sb + lines * lh + sa


def check(path, verbose=False):
    prs = Presentation(path)
    SW = prs.slide_width.inches
    SH = prs.slide_height.inches
    issues = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            try:
                L = sh.left.inches
                T = sh.top.inches
                W = sh.width.inches
                H = sh.height.inches
            except Exception:
                continue
            tag = '[s%d id%s %s]' % (si, sh.shape_id, sh.name)
            if L < -0.01 or T < -0.01 or L + W > SW + 0.01 or T + H > SH + 0.01:
                issues.append(('BOUNDS', tag, '形状越出页面',
                               'L=%.2f T=%.2f R=%.2f B=%.2f' % (L, T, L + W, T + H)))
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            txt = tf.text.strip()
            if not txt:
                continue
            ml = tf.margin_left.inches if tf.margin_left is not None else 0.1
            mr = tf.margin_right.inches if tf.margin_right is not None else 0.1
            mt = tf.margin_top.inches if tf.margin_top is not None else 0.05
            mb = tf.margin_bottom.inches if tf.margin_bottom is not None else 0.05
            w_pt = max(1.0, (W - ml - mr) * 72.0)
            avail = max(0.0, (H - mt - mb) * 72.0)
            need = sum(para_height(p, w_pt) for p in tf.paragraphs)
            if verbose:
                print('  %s avail=%.1f need=%.1f  %r' % (tag, avail, need, txt[:24]))
            if need > avail + 0.5:
                issues.append(('OVERFLOW', tag,
                               '文字超出框高 %.1fpt（可用 %.1f，需 %.1f）' % (need - avail, avail, need),
                               txt[:36].replace('\n', ' ')))
    return issues


if __name__ == '__main__':
    p = sys.argv[1]
    verbose = '-v' in sys.argv
    items = check(p, verbose)
    if not items:
        print('PASS：未发现溢出/越界')
    else:
        print('FAIL：发现 %d 处问题' % len(items))
        for kind, tag, msg, extra in items:
            print('  [%s] %s %s | %s' % (kind, tag, msg, extra))
